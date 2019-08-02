from bs4 import BeautifulSoup as bs 
import requests
from docx import Document

url = 'https://www.bastter.com'  # URL a ser raspada

page = requests.get(url)  # O retorno de um GET à URL
soup = bs(page.text, 'html.parser', from_encoding='utf-8')  # O esqueleto do HTML ficará aqui

todas_divs = soup.find_all('div', {'class': 'modal-body'})  # Captura todas as divs que tenham a classe "modal-body" em uma lista
div_regras = todas_divs[4]  # Pega a div do índice 4 da lista de divs, que é a que queremos
regras = div_regras.get_text()  # Pega somente o texto da div

document = Document()  # Cria o documento docx
document.add_heading('Regras Bastter.com')  # Adiciona um título para o documento
document.add_paragraph(regras)
document.save('bastter 2.docx')
